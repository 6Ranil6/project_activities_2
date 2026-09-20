from docx import Document
from docx.shared import Pt
from pathlib import Path
f=Path('Отчет_Хаялиев_ИУ6-74Б_Генерация_описаний_Авито.docx');d=Document(f)
toc=[p for p in d.paragraphs if '\t' in p.text and p.text.split('\t')[-1].isdigit()]
for i,p in enumerate(toc):
 p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_after=Pt(3);p.paragraph_format.page_break_before=(i==18)
refs=False;abbr=False
for p in d.paragraphs:
 if p.text=='ПЕРЕЧЕНЬ СОКРАЩЕНИЙ И ОБОЗНАЧЕНИЙ':abbr=True
 if p.text=='ВВЕДЕНИЕ':abbr=False
 if abbr:p.paragraph_format.line_spacing=1.5
 if p.text=='СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ':refs=True
 if p.text=='ПРИЛОЖЕНИЕ А':refs=False
 if refs:
  p.paragraph_format.line_spacing=1.5
  if p.text.startswith('7. '):p.paragraph_format.page_break_before=True
  if p.text.startswith('8. '):p.paragraph_format.page_break_before=False
  if p.text[:1].isdigit():p.paragraph_format.space_after=Pt(6)
d.save(f)

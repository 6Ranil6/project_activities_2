from docx import Document
from docx.shared import Cm
from copy import deepcopy
from pathlib import Path
f=Path('Отчет_Хаялиев_ИУ6-74Б_Генерация_описаний_Авито.docx');d=Document(f);a=Document('.report_work/restoration.docx')
for i in [111,127,176,213,239]:
 old=d.paragraphs[i]._p;old.getparent().replace(old,deepcopy(a.paragraphs[i+10]._p))
for i,file,width in [(101,'categories.png',16.5),(117,'biz_image2.png',14.5),(166,'slide_image1.png',16.5),(203,'rules.png',16.5),(229,'biz_image1.png',15)]:
 p=d.paragraphs[i];p.clear();p.alignment=1;p.paragraph_format.first_line_indent=0;p.paragraph_format.keep_with_next=True;p.add_run().add_picture('.report_work/'+file,width=Cm(width))
old=d.paragraphs[0]._p
for p in a.paragraphs[:11]:old.addprevious(deepcopy(p._p))
old.getparent().remove(old)
d.save(f)
